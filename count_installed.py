import openpyxl

wb = openpyxl.load_workbook('data/Thiết bị 2022.xlsx', data_only=True)

total_installed = 0
floor_counts = {}

for sheetname in wb.sheetnames:
    if not sheetname.startswith('Tầng'):
        continue
    sheet = wb[sheetname]
    count = 0
    # Room columns start at col 7
    rooms = []
    for c in range(7, sheet.max_column + 1):
        room_name = sheet.cell(4, c).value
        if room_name and str(room_name).strip() != 'Tổng':
            rooms.append((c, str(room_name).strip()))
    
    for r in range(5, sheet.max_row + 1):
        dev_name = sheet.cell(r, 3).value
        if not dev_name:
            continue
        for col_idx, room in rooms:
            val = sheet.cell(r, col_idx).value
            if val and isinstance(val, (int, float)) and val > 0:
                count += int(val)
    floor_counts[sheetname] = {'rooms_count': len(rooms), 'installed_devices_count': count}
    total_installed += count

with open('installed_counts.txt', 'w', encoding='utf-8') as f:
    f.write(f"Total installed devices across all floors: {total_installed}\n")
    for f_name, c_info in floor_counts.items():
        f.write(f"- {f_name}: {c_info['rooms_count']} phòng/vị trí, {c_info['installed_devices_count']} thiết bị đã lắp đặt\n")

print("Installed counts calculated.")
