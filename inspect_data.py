import openpyxl
import pypdf
import json

output = {}

# 1. Excel inspection
wb = openpyxl.load_workbook('data/Thiết bị 2022.xlsx', data_only=True)
output['excel_sheets'] = wb.sheetnames

excel_details = {}
for name in wb.sheetnames:
    sheet = wb[name]
    rows = []
    for r in range(1, min(30, sheet.max_row + 1)):
        row_vals = [sheet.cell(r, c).value for c in range(1, min(20, sheet.max_column + 1))]
        if any(v is not None for v in row_vals):
            rows.append({f"col_{c+1}": str(v) if v is not None else "" for c, v in enumerate(row_vals)})
    excel_details[name] = {
        'total_rows': sheet.max_row,
        'total_cols': sheet.max_column,
        'sample_rows': rows
    }
output['excel_details'] = excel_details

# 2. PDF inspection
reader = pypdf.PdfReader('data/SƠ ĐỒ CHUẨN.pdf')
output['pdf_pages'] = len(reader.pages)
pdf_text = []
for idx, page in enumerate(reader.pages):
    pdf_text.append({
        'page': idx + 1,
        'text': page.extract_text()
    })
output['pdf_details'] = pdf_text

with open('data_inspection.json', 'w', encoding='utf-8') as f:
    json.dump(output, f, ensure_ascii=False, indent=2)

print("Inspection completed. Written to data_inspection.json")
