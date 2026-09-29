import json
import openpyxl
import pypdf

with open('parse_output.txt', 'w', encoding='utf-8') as out:
    # 1. Inspect PDF Text
    reader = pypdf.PdfReader('data/SƠ ĐỒ CHUẨN.pdf')
    out.write(f"Total PDF pages: {len(reader.pages)}\n")
    for i, page in enumerate(reader.pages):
        out.write(f"\n================ PAGE {i+1} ================\n")
        text = page.extract_text()
        out.write(text if text else "[EMPTY TEXT / SCANNED]\n")

    # 2. Inspect Excel sheets
    wb = openpyxl.load_workbook('data/Thiết bị 2022.xlsx', data_only=True)
    for sheetname in wb.sheetnames:
        sheet = wb[sheetname]
        out.write(f"\n================ SHEET: {sheetname} (max_row={sheet.max_row}, max_col={sheet.max_column}) ================\n")
        for r in range(1, min(35, sheet.max_row + 1)):
            row_vals = [sheet.cell(r, c).value for c in range(1, min(25, sheet.max_column + 1))]
            if any(v is not None for v in row_vals):
                out.write(f"R{r}: {row_vals}\n")
