import openpyxl
import json

wb = openpyxl.load_workbook('data/Thiết bị 2022.xlsx', data_only=True)

# Sheet TOTAL
total_sheet = wb['TOTAL']
total_devices = []
for r in range(5, total_sheet.max_row + 1):
    stt = total_sheet.cell(r, 1).value
    cat = total_sheet.cell(r, 2).value
    name = total_sheet.cell(r, 3).value
    code = total_sheet.cell(r, 4).value
    spec = total_sheet.cell(r, 5).value
    mfg = total_sheet.cell(r, 6).value
    qty_in = total_sheet.cell(r, 7).value
    qty_real = total_sheet.cell(r, 8).value
    diff = total_sheet.cell(r, 9).value

    if name:
        total_devices.append({
            'stt': stt,
            'category': str(cat).strip() if cat else '',
            'name': str(name).strip() if name else '',
            'model': str(code).strip() if code else '',
            'spec': str(spec).strip() if spec else '',
            'mfg': str(mfg).strip() if mfg else '',
            'qty_in': qty_in,
            'qty_real': qty_real
        })

print(f"Total sheet devices count: {len(total_devices)}")

with open('total_devices_extracted.json', 'w', encoding='utf-8') as f:
    json.dump(total_devices, f, ensure_ascii=False, indent=2)

print("Saved to total_devices_extracted.json successfully")
